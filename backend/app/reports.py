from io import BytesIO


def _dinero(valor: float) -> str:
    return f"S/ {valor:,.2f}"


def generar_pdf(lote, cierre, panel, registros, vacunas, tratamientos, alertas) -> BytesIO:
    from reportlab.graphics.shapes import Drawing, PolyLine, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    salida = BytesIO()
    doc = SimpleDocTemplate(
        salida,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=f"Cierre de lote {lote.id} - AviSmart",
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph("AviSmart · Reporte de cierre de lote", estilos["Title"]),
        Paragraph(
            f"{lote.galpon.granja.nombre} · {lote.galpon.nombre} · línea {lote.linea}",
            estilos["Heading2"],
        ),
        Spacer(1, 5 * mm),
    ]

    resumen = [
        ["Ingreso", lote.fecha_ingreso.isoformat(), "Cierre", cierre.fecha_cierre.isoformat()],
        ["Población inicial", lote.poblacion_inicial, "Aves vendidas", cierre.aves_vendidas],
        ["Mortalidad", f'{panel["mortalidad_acumulada"]:.2f}%', "FCR", panel["fcr"] or "—"],
        ["Peso promedio de venta", f'{cierre.peso_total_kg * 1000 / cierre.aves_vendidas:,.0f} g', "EPEF", panel["epef"] or "—"],
        ["Último pesaje", f'{panel["peso_actual_g"]:.0f} g', "Peso total vendido", f"{cierre.peso_total_kg:,.1f} kg"],
    ]
    tabla = Table(resumen, colWidths=[38 * mm, 48 * mm, 38 * mm, 48 * mm])
    tabla.setStyle(_estilo_tabla())
    elementos.extend([tabla, Spacer(1, 5 * mm), Paragraph("Recepción de pollitos", estilos["Heading2"])])
    filas = _recepcion_filas(lote)
    tabla = Table(filas, colWidths=[38 * mm, 48 * mm, 38 * mm, 48 * mm])
    tabla.setStyle(_estilo_tabla())
    tabla.setStyle([("SPAN", (1, len(filas) - 1), (3, len(filas) - 1))])  # la fila de evidencias ocupa todo el ancho
    elementos.extend([tabla])
    if lote.observaciones:
        elementos.append(Paragraph(f"Observaciones de la recepción: {lote.observaciones}", estilos["BodyText"]))
    elementos.extend([Spacer(1, 6 * mm), Paragraph("Evolución de peso", estilos["Heading2"])])

    puntos = [(i, r.peso_muestra_g) for i, r in enumerate(registros)]
    if len(puntos) >= 2:
        ancho, alto = 470, 150
        maximo = max(peso for _, peso in puntos) or 1
        escala_x = ancho / max(len(puntos) - 1, 1)
        escala_y = (alto - 25) / maximo
        dibujo = Drawing(ancho, alto)
        dibujo.add(String(0, alto - 12, "Peso de muestra (g)", fontSize=8))
        dibujo.add(
            PolyLine(
                [(i * escala_x, 10 + peso * escala_y) for i, peso in puntos],
                strokeColor=colors.HexColor("#35624c"),
                strokeWidth=2,
            )
        )
        elementos.append(dibujo)
    else:
        elementos.append(Paragraph("No hay suficientes pesajes para trazar el gráfico.", estilos["BodyText"]))

    elementos.extend([Spacer(1, 5 * mm), Paragraph("Resultado económico", estilos["Heading2"])])
    economia = [
        ["Ingreso por venta", _dinero(cierre.ingreso_total)],
        ["Costo de pollitos", _dinero(cierre.costo_pollitos)],
        ["Costo de alimento", _dinero(cierre.costo_alimento)],
        ["Costo sanitario", _dinero(cierre.costo_sanidad)],
        ["Otros costos", _dinero(cierre.otros_costos)],
        ["Costos totales", _dinero(cierre.total_costos)],
        ["Margen", _dinero(cierre.margen)],
    ]
    tabla = Table(economia, colWidths=[80 * mm, 55 * mm])
    tabla.setStyle(_estilo_tabla())
    elementos.extend([tabla, PageBreak(), Paragraph("Registros diarios", estilos["Heading2"])])

    filas = [["Fecha", "Muertes", "Alimento kg", "Agua L", "Peso g"]]
    filas.extend(
        [[r.fecha.isoformat(), r.mortalidad, r.alimento_kg, r.agua_l, r.peso_muestra_g] for r in registros]
    )
    tabla = Table(filas, repeatRows=1, colWidths=[35 * mm, 28 * mm, 35 * mm, 30 * mm, 30 * mm])
    tabla.setStyle(_estilo_tabla(encabezado=True))
    elementos.extend([tabla, Spacer(1, 6 * mm), Paragraph("Resumen sanitario", estilos["Heading2"])])

    sanitario = [["Tipo", "Nombre", "Fecha", "Estado"]]
    sanitario.extend([["Vacuna", v.nombre, v.fecha_programada.isoformat(), v.estado] for v in vacunas])
    sanitario.extend([["Tratamiento", t.medicamento, t.fecha_inicio.isoformat(), t.estado] for t in tratamientos])
    tabla = Table(sanitario, repeatRows=1, colWidths=[32 * mm, 55 * mm, 35 * mm, 35 * mm])
    tabla.setStyle(_estilo_tabla(encabezado=True))
    elementos.extend(
        [
            tabla,
            Spacer(1, 5 * mm),
            Paragraph(f"Alertas registradas: {len(alertas)}", estilos["BodyText"]),
            Paragraph(f"Observaciones de cierre: {cierre.observaciones or 'Sin observaciones.'}", estilos["BodyText"]),
        ]
    )
    doc.build(elementos)
    salida.seek(0)
    return salida


def _o(valor, sufijo: str = "") -> str:
    return "—" if valor in (None, "") else f"{valor}{sufijo}"


SEXOS = {"MIXTO": "Mixto", "MACHOS": "Machos", "HEMBRAS": "Hembras"}


def _recepcion_filas(lote) -> list[list]:
    """Datos del ingreso del lote; los lotes antiguos pueden no tenerlos."""
    costo = None
    if lote.precio_pollito is not None:
        costo = _dinero(lote.precio_pollito * (lote.cantidad_recibida or lote.poblacion_inicial))
    return [
        ["Incubadora / proveedor", _o(lote.proveedor), "Guía de remisión", _o(lote.guia)],
        ["Sexo", SEXOS.get(lote.sexo or "", _o(lote.sexo)), "Fecha de llegada", lote.fecha_ingreso.isoformat()],
        ["Pollitos recibidos", _o(lote.cantidad_recibida), "Descartados al llegar", _o(lote.descarte_llegada)],
        ["Alojados en el galpón", lote.poblacion_inicial, "Peso al llegar", _o(lote.peso_llegada_g, " g")],
        ["Precio por pollito", _dinero(lote.precio_pollito) if lote.precio_pollito is not None else "—", "Costo de pollitos", _o(costo)],
        ["Meta de peso", _o(lote.peso_objetivo_g, " g"), "Edad de saca prevista", _o(lote.edad_saca_dias, " días")],
        ["Evidencias adjuntas", _evidencias(lote), "", ""],
    ]


def _evidencias(lote) -> str:
    """Resumen de las evidencias de la apertura (guía, certificados, fotos)."""
    from sqlalchemy.orm import object_session

    from .evidencias import resumen

    sesion = object_session(lote)
    filas = resumen(sesion, lote.id) if sesion is not None else []
    return "; ".join(f"{nombre} ({cantidad})" for nombre, cantidad in filas) or "—"


def _estilo_tabla(encabezado: bool = False):
    from reportlab.lib import colors
    from reportlab.platypus import TableStyle

    reglas = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5ce")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#f4f7f5")]),
    ]
    if encabezado:
        reglas.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#35624c")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]
        )
    return TableStyle(reglas)


def generar_excel(lote, cierre, panel, registros, vacunas, tratamientos, alertas) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill

    libro = Workbook()
    resumen = libro.active
    resumen.title = "Cierre"
    resumen.append(["AviSmart", "Reporte de cierre de lote"])
    resumen.append(["Granja", lote.galpon.granja.nombre])
    resumen.append(["Galpón", lote.galpon.nombre])
    resumen.append(["Línea", lote.linea])
    resumen.append(["Fecha de ingreso", lote.fecha_ingreso])
    resumen.append(["Fecha de cierre", cierre.fecha_cierre])
    resumen.append(["Población inicial", lote.poblacion_inicial])
    resumen.append(["Aves vendidas", cierre.aves_vendidas])
    resumen.append(["Peso total vendido (kg)", cierre.peso_total_kg])
    resumen.append(["Peso promedio de venta (g)", round(cierre.peso_total_kg * 1000 / cierre.aves_vendidas)])
    resumen.append(["Mortalidad acumulada (%)", panel["mortalidad_acumulada"]])
    resumen.append(["FCR", panel["fcr"]])
    resumen.append(["EPEF", panel["epef"]])
    resumen.append(["Ingreso total", cierre.ingreso_total])
    resumen.append(["Costos totales", cierre.total_costos])
    resumen.append(["Margen", cierre.margen])
    resumen.append(["Observaciones", cierre.observaciones])

    recepcion = libro.create_sheet("Recepción", 1)
    recepcion.append(["Dato", "Valor"])
    for fila in [
        ["Incubadora / proveedor", lote.proveedor],
        ["Guía de remisión", lote.guia],
        ["Sexo", SEXOS.get(lote.sexo or "", lote.sexo)],
        ["Fecha de llegada", lote.fecha_ingreso],
        ["Pollitos recibidos", lote.cantidad_recibida],
        ["Descartados al llegar", lote.descarte_llegada],
        ["Alojados en el galpón", lote.poblacion_inicial],
        ["Peso al llegar (g)", lote.peso_llegada_g],
        ["Precio por pollito (S/)", lote.precio_pollito],
        ["Costo de pollitos (S/)", round(lote.precio_pollito * (lote.cantidad_recibida or lote.poblacion_inicial), 2) if lote.precio_pollito is not None else None],
        ["Meta de peso (g)", lote.peso_objetivo_g],
        ["Edad de saca prevista (días)", lote.edad_saca_dias],
        ["Observaciones de la recepción", lote.observaciones],
    ]:
        recepcion.append(fila)

    hoja = libro.create_sheet("Registros diarios")
    hoja.append(["Fecha", "Mortalidad", "Alimento kg", "Agua L", "Peso muestra g", "Tipo de alimento"])
    for r in registros:
        hoja.append([r.fecha, r.mortalidad, r.alimento_kg, r.agua_l, r.peso_muestra_g, r.fase_alimento or ""])
    if registros:
        grafico = LineChart()
        grafico.title = "Evolución del peso de muestra"
        grafico.y_axis.title = "Peso (g)"
        grafico.x_axis.title = "Fecha"
        datos = Reference(hoja, min_col=5, min_row=1, max_row=1 + len(registros))
        categorias = Reference(hoja, min_col=1, min_row=2, max_row=1 + len(registros))
        grafico.add_data(datos, titles_from_data=True)
        grafico.set_categories(categorias)
        hoja.add_chart(grafico, "H2")

    sanidad = libro.create_sheet("Sanidad")
    sanidad.append(["Tipo", "Nombre", "Fecha programada / inicio", "Aplicada", "Estado", "Detalle"])
    for v in vacunas:
        sanidad.append(["Vacuna", v.nombre, v.fecha_programada, v.fecha_aplicacion, v.estado, v.observacion])
    for t in tratamientos:
        sanidad.append(["Tratamiento", t.medicamento, t.fecha_inicio, None, t.estado, t.diagnostico])

    hoja_alertas = libro.create_sheet("Alertas")
    hoja_alertas.append(["Fecha", "Tipo", "Mensaje", "Estado", "Observación"])
    for a in alertas:
        hoja_alertas.append([a.creada_en, a.tipo, a.mensaje, a.estado, a.observacion])

    for ws in libro.worksheets:
        ws.freeze_panes = "A2"
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="35624C")
        for columna in ws.columns:
            letra = columna[0].column_letter
            ancho = min(max(len(str(c.value or "")) for c in columna) + 2, 55)
            ws.column_dimensions[letra].width = ancho

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida
