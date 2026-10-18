"""Límite de intentos del login, recepción en los reportes, aviso de solicitudes y Telegram por cliente."""
import os
import unittest
from datetime import date, datetime, timedelta
from io import BytesIO
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db import Base, get_db
from app.main import app
from app.models import CierreLote, Galpon, Granja, IntentoLogin, Lote, Organizacion, RegistroDiario, SolicitudContacto, Usuario
from app.security import hash_password

CLAVE = "ClavePrueba123"
HOY = date.today()


@patch.dict(os.environ, {"AVISMART_PROXIES_CONFIABLES": "testclient", "SMTP_HOST": "", "AVISMART_DEV": "true"})
class Punto3Test(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        self.Session = sessionmaker(bind=self.engine, autoflush=False, autocommit=False)
        Base.metadata.create_all(self.engine)

        def db_prueba():
            db = self.Session()
            try:
                yield db
            finally:
                db.close()

        app.dependency_overrides[get_db] = db_prueba
        db = self.Session()
        org = Organizacion(nombre="Avícola Andina", estado="ACTIVA", creada_en=datetime.now())
        db.add(org)
        db.flush()
        self.org_id = org.id
        db.add_all([
            Usuario(nombre="Ana", email="ana@granja.pe", password_hash=hash_password(CLAVE), rol="PRODUCTOR", activo=1, debe_cambiar_clave=0, organizacion_id=org.id),
            Usuario(nombre="Óscar", email="oscar@granja.pe", password_hash=hash_password(CLAVE), rol="OPERARIO", activo=1, debe_cambiar_clave=0, organizacion_id=org.id),
            Usuario(nombre="Dueño", email="duenio@avismart.pe", password_hash=hash_password(CLAVE), rol="PLATAFORMA", activo=1, debe_cambiar_clave=0),
        ])
        granja = Granja(nombre="Granja Andina", ubicacion="Juliaca", organizacion_id=org.id)
        db.add(granja)
        db.flush()
        galpon = Galpon(granja_id=granja.id, nombre="Galpón A", capacidad=5000)
        db.add(galpon)
        db.flush()
        lote = Lote(galpon_id=galpon.id, linea="Cobb", fecha_ingreso=HOY - timedelta(days=42), poblacion_inicial=4970, peso_objetivo_g=2800, estado="CERRADO",
                    fecha_cierre=HOY, proveedor="Incubadora San Fernando", guia="T001-000123", sexo="MIXTO", cantidad_recibida=5000,
                    descarte_llegada=30, peso_llegada_g=42.5, precio_pollito=2.1, edad_saca_dias=42, observaciones="Llegaron activos")
        db.add(lote)
        db.flush()
        self.lote_id = lote.id
        db.add(RegistroDiario(lote_id=lote.id, usuario_id=1, fecha=HOY - timedelta(days=1), mortalidad=3, alimento_kg=900, agua_l=1800, peso_muestra_g=2900))
        db.add(CierreLote(lote_id=lote.id, fecha_cierre=HOY, aves_vendidas=4800, peso_total_kg=13900, precio_venta_kg=7.5, ingreso_total=104250,
                          costo_pollitos=10500, costo_alimento=40000, costo_sanidad=900, otros_costos=0, total_costos=51400, margen=52850,
                          observaciones="", cerrado_por=1, created_at=datetime.now()))
        db.commit()
        db.close()
        self.client = TestClient(app)

    def tearDown(self):
        app.dependency_overrides.clear()
        self.engine.dispose()

    def login(self, email, clave, ip="10.9.0.1"):
        return self.client.post("/api/auth/login", json={"email": email, "password": clave}, headers={"X-Forwarded-For": ip})

    def h(self, email):
        return {"Authorization": "Bearer " + self.login(email, CLAVE, ip="10.0.0.99").json()["token"]}

    # ------------------------------------------------------------ login
    def test_login_se_bloquea_tras_5_fallos_y_avisa_antes(self):
        for i in range(4):
            r = self.login("ana@granja.pe", "mala")
            self.assertEqual(r.status_code, 401)
        self.assertIn("Te queda 1 intento", r.json()["detail"])
        self.assertEqual(self.login("ana@granja.pe", "mala").status_code, 401)
        bloqueado = self.login("ana@granja.pe", CLAVE)
        self.assertEqual(bloqueado.status_code, 429, "ni con la clave correcta mientras dura el bloqueo")
        self.assertIn("Demasiados intentos", bloqueado.json()["detail"])
        self.assertEqual(self.login("oscar@granja.pe", CLAVE, ip="10.9.0.2").status_code, 200, "otras cuentas siguen entrando")

    def test_bloqueo_vence_y_un_acierto_limpia_los_fallos(self):
        for _ in range(5):
            self.login("ana@granja.pe", "mala")
        db = self.Session()
        db.query(IntentoLogin).update({"creado_en": datetime.now() - timedelta(minutes=16)})
        db.commit()
        db.close()
        self.assertEqual(self.login("ana@granja.pe", CLAVE).status_code, 200)
        db = self.Session()
        self.assertEqual(db.query(IntentoLogin).filter_by(email="ana@granja.pe").count(), 0)
        db.close()

    def test_correo_inexistente_recibe_el_mismo_trato(self):
        for _ in range(5):
            self.assertEqual(self.login("nadie@granja.pe", "x").status_code, 401)
        self.assertEqual(self.login("nadie@granja.pe", "x").status_code, 429)

    def test_limite_por_ip(self):
        for i in range(20):
            self.login(f"x{i}@granja.pe", "mala", ip="200.0.0.1")
        self.assertEqual(self.login("ana@granja.pe", CLAVE, ip="200.0.0.1").status_code, 429)
        self.assertEqual(self.login("ana@granja.pe", CLAVE, ip="200.0.0.2").status_code, 200)

    # ------------------------------------------------------------ reportes
    def test_reportes_incluyen_la_recepcion(self):
        excel = self.client.get(f"/api/lotes/{self.lote_id}/reporte.xlsx", headers=self.h("ana@granja.pe"))
        hoja = load_workbook(BytesIO(excel.content))["Recepción"]
        datos = {fila[0].value: fila[1].value for fila in hoja.iter_rows(min_row=2)}
        self.assertEqual(datos["Incubadora / proveedor"], "Incubadora San Fernando")
        self.assertEqual(datos["Guía de remisión"], "T001-000123")
        self.assertEqual(datos["Descartados al llegar"], 30)
        self.assertEqual(datos["Peso al llegar (g)"], 42.5)
        self.assertEqual(datos["Costo de pollitos (S/)"], 10500)
        pdf = self.client.get(f"/api/lotes/{self.lote_id}/reporte.pdf", headers=self.h("ana@granja.pe"))
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    # ------------------------------------------------------------ aviso de solicitud
    def test_solicitud_confirmada_avisa_a_la_plataforma(self):
        db = self.Session()
        from app.main import _hash_token

        db.add(SolicitudContacto(nombre="Luis", email="luis@sur.pe", empresa="Avícola del Sur", pais="Perú", aves="", motivo="COMERCIAL",
                                 plan="BASICO", mensaje="Quiero ordenar mis galpones", estado="PENDIENTE", creada_en=datetime.now(),
                                 token_hash=_hash_token("t" * 30), token_expira=datetime.now() + timedelta(hours=1)))
        db.commit()
        db.close()
        with patch("app.main.notifications.avisar_plataforma") as telegram, patch("app.main.correo.enviar_aviso_solicitud") as mail:
            r = self.client.post("/api/contacto/confirmar", json={"token": "t" * 30})
        self.assertEqual(r.status_code, 200)
        self.assertIn("Avícola del Sur", telegram.call_args.args[0])
        mail.assert_called_once()
        self.assertEqual(mail.call_args.args[0], "duenio@avismart.pe")

    # ------------------------------------------------------------ Telegram por cliente
    @patch.dict(os.environ, {"TELEGRAM_BOT_TOKEN": "token-prueba"})
    def test_vincular_chat_del_cliente_con_codigo(self):
        prod = self.h("ana@granja.pe")
        self.assertEqual(self.client.post("/api/organizacion/telegram/vincular", headers=self.h("oscar@granja.pe")).status_code, 403)

        def telegram(metodo, datos=None):
            if metodo == "getMe":
                return True, {"username": "AviSmartBot"}
            if metodo == "getUpdates":
                return True, [
                    {"message": {"text": "/start AVIOTRO123", "chat": {"id": 111, "title": "Otro cliente"}}},
                    {"message": {"text": f"/start {codigo['codigo']}", "chat": {"id": -100555, "title": "Galponeros Andina"}}},
                ]
            return True, {}

        from app import notifications

        notifications._usuario_bot_cache.cache_clear()
        with patch("app.notifications._llamar", side_effect=telegram):
            codigo = self.client.post("/api/organizacion/telegram/vincular", headers=prod).json()
            self.assertTrue(codigo["codigo"].startswith("AVI"))
            self.assertEqual(codigo["enlace"], f"https://t.me/AviSmartBot?start={codigo['codigo']}")
            r = self.client.post("/api/organizacion/telegram/verificar", headers=prod)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["vinculado"])
        self.assertEqual(r.json()["chat_nombre"], "Galponeros Andina", "toma el chat que envió SU código, no otro")
        db = self.Session()
        self.assertEqual(db.get(Organizacion, self.org_id).telegram_chat_id, "-100555")
        db.close()
        notifications._usuario_bot_cache.cache_clear()

    @patch.dict(os.environ, {"TELEGRAM_BOT_TOKEN": "token-prueba"})
    def test_verificar_sin_codigo_enviado(self):
        prod = self.h("ana@granja.pe")
        from app import notifications

        notifications._usuario_bot_cache.cache_clear()
        with patch("app.notifications._llamar", side_effect=lambda m, d=None: (True, {"username": "AviSmartBot"} if m == "getMe" else [])):
            self.client.post("/api/organizacion/telegram/vincular", headers=prod)
            r = self.client.post("/api/organizacion/telegram/verificar", headers=prod)
        self.assertEqual(r.status_code, 400)
        self.assertIn("Todavía no vemos el código", r.json()["detail"])
        notifications._usuario_bot_cache.cache_clear()


if __name__ == "__main__":
    unittest.main()
