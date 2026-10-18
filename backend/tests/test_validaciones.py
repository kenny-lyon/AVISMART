"""Validaciones de datos, manejo de errores del servidor y reglas del cierre de lote."""
import unittest
from datetime import date, datetime, timedelta
from unittest.mock import patch

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.db import Base, get_db
from app.logic import peso_esperado
from app.main import app
from app.models import Galpon, Granja, Lote, Organizacion, RegistroDiario, Usuario
from app.security import hash_password

HOY = date.today()
CLAVE = "ClavePrueba123"


class ValidacionesTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        self.Session = sessionmaker(bind=self.engine, autoflush=False, autocommit=False)
        Base.metadata.create_all(self.engine)

        def db_prueba():
            db = self.Session()
            try:
                yield db
            finally:
                db.close()

        app.dependency_overrides[get_db] = db_prueba
        db = self.Session()
        org = Organizacion(nombre="Cliente", estado="ACTIVA", creada_en=datetime.now())
        db.add(org)
        db.flush()
        for rol in ("PRODUCTOR", "OPERARIO"):
            db.add(Usuario(nombre=rol.title(), email=f"{rol.lower()}@granja.pe", password_hash=hash_password(CLAVE), rol=rol, activo=1, debe_cambiar_clave=0, organizacion_id=org.id))
        granja = Granja(nombre="Granja", ubicacion="Puno", organizacion_id=org.id)
        db.add(granja)
        db.flush()
        galpon = Galpon(granja_id=granja.id, nombre="Galpón 1", capacidad=2000)
        db.add(galpon)
        db.flush()
        self.ingreso = HOY - timedelta(days=30)
        lote = Lote(galpon_id=galpon.id, linea="Ross", fecha_ingreso=self.ingreso, poblacion_inicial=1000, peso_objetivo_g=2800, estado="ACTIVO")
        db.add(lote)
        db.commit()
        self.lote_id = lote.id
        db.close()
        self.client = TestClient(app, raise_server_exceptions=False)
        self.h = {rol: self.entrar(f"{rol.lower()}@granja.pe") for rol in ("PRODUCTOR", "OPERARIO")}

    def tearDown(self):
        app.dependency_overrides.clear()
        self.engine.dispose()

    def entrar(self, email):
        return {"Authorization": "Bearer " + self.client.post("/api/auth/login", json={"email": email, "password": CLAVE}).json()["token"]}

    def registro(self, dias_atras=1, **cambios):
        edad = 30 - dias_atras
        datos = {"fecha": (HOY - timedelta(days=dias_atras)).isoformat(), "mortalidad": 3, "alimento_kg": 120, "agua_l": 230,
                 "peso_muestra_g": round(peso_esperado(edad, "Ross"))}
        datos.update(cambios)
        return self.client.post(f"/api/lotes/{self.lote_id}/registros", headers=self.h["OPERARIO"], json=datos)

    def cerrar(self, fecha=HOY, **cambios):
        datos = {"fecha_cierre": fecha.isoformat(), "aves_vendidas": 990, "peso_total_kg": 1900, "precio_venta_kg": 7.5,
                 "costo_pollitos": 2100, "costo_alimento": 4000, "costo_sanidad": 150, "otros_costos": 0}
        datos.update(cambios)
        return self.client.post(f"/api/lotes/{self.lote_id}/cerrar", headers=self.h["PRODUCTOR"], json=datos)

    # ------------------------------------------------------------ registro diario
    def test_registro_valido(self):
        self.assertEqual(self.registro().status_code, 200)

    def test_registro_rechaza_datos_absurdos(self):
        peso = self.registro(peso_muestra_g=99999)
        self.assertEqual(peso.status_code, 422)
        self.assertIn("peso por ave (g)", peso.json()["detail"])
        imposible = self.registro(peso_muestra_g=5000)
        self.assertEqual(imposible.status_code, 400)
        self.assertIn("no es posible para un pollo de 29 días", imposible.json()["detail"])
        self.assertIn("350 g por ave", self.registro(alimento_kg=900).json()["detail"])
        self.assertIn("800 mL por ave", self.registro(agua_l=2000).json()["detail"])
        self.assertIn("superan las aves vivas", self.registro(mortalidad=1500).json()["detail"])

    def test_dos_registros_simultaneos_no_dan_error_500(self):
        error = IntegrityError("INSERT", {}, Exception("UNIQUE constraint failed"))
        with patch.object(Session, "flush", side_effect=error):
            r = self.registro()
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["detail"], "Ya hay un registro para esa fecha")

    # ------------------------------------------------------------ clima y sanidad
    def test_clima_fuera_de_rango_con_mensaje_en_espanol(self):
        r = self.client.post(f"/api/lotes/{self.lote_id}/clima", headers=self.h["OPERARIO"], json={"temperatura": 500, "humedad": 60})
        self.assertEqual(r.status_code, 422)
        self.assertEqual(r.json()["detail"], "«temperatura (°C)» debe ser menor o igual a 50.")

    def test_tratamiento_futuro_o_con_mas_aves_que_las_vivas(self):
        base = {"medicamento": "Enrofloxacina 10 %", "dosis": "10 mg/kg", "via": "Agua de bebida", "dias_tratamiento": 3, "dias_carencia": 7}
        ruta = f"/api/lotes/{self.lote_id}/tratamientos"
        futuro = self.client.post(ruta, headers=self.h["OPERARIO"], json={**base, "fecha_inicio": (HOY + timedelta(days=365)).isoformat()})
        self.assertEqual(futuro.status_code, 400)
        self.assertIn("no puede ser futura", futuro.json()["detail"])
        muchas = self.client.post(ruta, headers=self.h["OPERARIO"], json={**base, "fecha_inicio": HOY.isoformat(), "aves_tratadas": 1500})
        self.assertIn("superan las aves vivas", muchas.json()["detail"])

    def test_vacuna_fuera_de_la_campana(self):
        ruta = f"/api/lotes/{self.lote_id}/sanidad"
        for fecha in (self.ingreso - timedelta(days=1), self.ingreso + timedelta(days=71), date(2000, 1, 1)):
            r = self.client.post(ruta, headers=self.h["OPERARIO"], json={"nombre": "Marek (HVT)", "fecha_programada": fecha.isoformat()})
            self.assertEqual(r.status_code, 400, fecha)
        ok = self.client.post(ruta, headers=self.h["OPERARIO"], json={"nombre": "Marek (HVT)", "fecha_programada": self.ingreso.isoformat()})
        self.assertEqual(ok.status_code, 200)

    # ------------------------------------------------------------ cierre
    def test_no_se_cierra_sin_registros(self):
        r = self.cerrar()
        self.assertEqual(r.status_code, 400)
        self.assertIn("no tiene registros diarios", r.json()["detail"])

    def test_fecha_de_faena_elegible_y_coherente(self):
        self.registro(dias_atras=5)
        anterior = self.cerrar(HOY - timedelta(days=6))
        self.assertIn("la fecha de faena no puede ser anterior", anterior.json()["detail"])
        absurdo = self.cerrar(HOY - timedelta(days=3), peso_total_kg=99999)
        self.assertIn("peso promedio sería", absurdo.json()["detail"])
        r = self.cerrar(HOY - timedelta(days=3))
        self.assertEqual(r.status_code, 200, r.text)
        lote = next(l for l in self.client.get("/api/resumen", headers=self.h["PRODUCTOR"]).json()["lotes"] if l["id"] == self.lote_id)
        self.assertEqual(lote["fecha_cierre"], (HOY - timedelta(days=3)).isoformat())
        self.assertEqual(lote["dia_vida"], 27, "la edad del lote se congela en la fecha de faena")

    def test_reporte_muestra_peso_promedio_de_venta(self):
        self.registro()
        self.assertEqual(self.cerrar().status_code, 200)
        from io import BytesIO

        from openpyxl import load_workbook

        excel = self.client.get(f"/api/lotes/{self.lote_id}/reporte.xlsx", headers=self.h["PRODUCTOR"])
        hoja = load_workbook(BytesIO(excel.content))["Cierre"]
        filas = {fila[0].value: fila[1].value for fila in hoja.iter_rows() if fila[0].value}
        self.assertEqual(filas["Peso promedio de venta (g)"], round(1900 * 1000 / 990))

    def test_error_inesperado_responde_json_claro(self):
        with patch("app.main.indicadores", side_effect=RuntimeError("falla interna")):
            r = self.client.get(f"/api/lotes/{self.lote_id}/panel", headers=self.h["PRODUCTOR"])
        self.assertEqual(r.status_code, 500)
        self.assertIn("error inesperado", r.json()["detail"])


if __name__ == "__main__":
    unittest.main()
